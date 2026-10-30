"""
Unit & Integration Tests for Milestone 3:
Multi-Format Document Ingestion & Bilingual Layout-Preserving Translation Subsystems.
"""

from __future__ import annotations

import io
import docx
import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.ai.models import TranslationRequest
from app.ai.translation import (
    DOMAIN_GLOSSARIES,
    OFFLINE_PHRASE_MAPPINGS,
    TranslationMemory,
    execute_document_translation,
    split_text_into_segments,
    translate_segment_deterministic,
)
from app.services.multi_format import (
    detect_document_format,
    parse_csv_to_structure,
    parse_docx_to_structure,
    parse_markdown_to_structure,
    synthesize_pdf_from_structure,
)


@pytest.fixture
def client():
    return TestClient(app)


def test_detect_document_format():
    assert detect_document_format("test.pdf", b"%PDF-1.7\n...") == "pdf"
    assert detect_document_format("doc.docx", b"PK\x03\x04...") == "docx"
    assert detect_document_format("sheet.xlsx", b"PK\x03\x04...") in ("xlsx", "docx")
    assert detect_document_format("data.csv", b"a,b,c\n1,2,3") == "csv"
    assert detect_document_format("notes.md", b"# Header\n- item 1") == "markdown"
    assert detect_document_format("plain.txt", b"plain text line") == "text"


def test_parse_markdown_to_structure():
    md = b"""# Corporate Governance Report
This report outlines enterprise SOC 2 and ISO 27001 compliance.

## Audit Summary
- Identity and Access Management
- Zero Trust Network Architecture

| Control | Status | Evidence |
|---|---|---|
| IAM-01 | Verified | Okta / MFA |
| NET-04 | Passed | TLS 1.3 / mTLS |
"""
    struct = parse_markdown_to_structure(md)
    assert struct["format"] == "markdown"
    assert struct["title"] == "Corporate Governance Report"
    assert struct["table_count"] == 1
    assert any(b["type"] == "heading" for b in struct["blocks"])
    assert any(b["type"] == "list_item" for b in struct["blocks"])
    assert any(b["type"] == "table" for b in struct["blocks"])


def test_parse_csv_to_structure():
    csv_bytes = b"Quarter,Revenue,OperatingMargin\nQ1,$12M,24%\nQ2,$14.5M,26%\nQ3,$16M,28%\n"
    struct = parse_csv_to_structure(csv_bytes)
    assert struct["format"] == "csv"
    assert struct["table_count"] == 1
    table_block = [b for b in struct["blocks"] if b["type"] == "table"][0]
    assert len(table_block["rows"]) == 4  # Header + 3 data rows
    assert table_block["rows"][0] == ["Quarter", "Revenue", "OperatingMargin"]


def test_parse_docx_to_structure():
    doc = docx.Document()
    doc.add_heading("Contract Amendment Agreement", level=1)
    doc.add_paragraph("This amendment governs intellectual property transfer.")
    tbl = doc.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "Party"
    tbl.cell(0, 1).text = "Signatory"
    tbl.cell(1, 0).text = "Acme Corp"
    tbl.cell(1, 1).text = "Jane Doe, CEO"

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    struct = parse_docx_to_structure(docx_bytes)
    assert struct["format"] == "docx"
    assert "Contract Amendment" in struct["title"]
    assert struct["table_count"] == 1
    assert any(b["type"] == "paragraph" for b in struct["blocks"])


def test_synthesize_pdf_from_structure():
    struct = {
        "format": "markdown",
        "title": "Synthesized Test Document",
        "blocks": [
            {"type": "heading", "level": 1, "text": "Section 1: Architecture"},
            {"type": "paragraph", "text": "This is a high-fidelity synthesized PDF rendering test."},
            {"type": "list_item", "text": "Item A: Multi-Tenant Database Isolation"},
            {"type": "list_item", "text": "Item B: Local-First Privacy Protection"},
            {
                "type": "table",
                "rows": [
                    ["Metric", "Target", "Achieved"],
                    ["P95 Latency", "< 200ms", "45ms"],
                    ["Recall@3", ">= 80%", "92%"],
                ],
            },
        ],
        "table_count": 1,
    }
    pdf_bytes = synthesize_pdf_from_structure(struct, filename="test_doc.md")
    assert pdf_bytes.startswith(b"%PDF-")
    assert len(pdf_bytes) > 1000


def test_split_text_into_segments():
    text = """# Section 1
First paragraph here.
Second sentence in paragraph.

- First bullet
- Second bullet

| Col 1 | Col 2 |
|---|---|
| A | B |
"""
    segments = split_text_into_segments(text)
    types = [s["type"] for s in segments]
    assert "heading" in types
    assert "paragraph" in types
    assert "list_item" in types
    assert "table_cell" in types


def test_translation_memory_caching():
    mem = TranslationMemory(max_size=100)
    mem.put("en", "es", "legal", "Force Majeure", "fuerza mayor", "local_deterministic", ["force majeure"])
    cached = mem.get("en", "es", "legal", "Force Majeure")
    assert cached is not None
    assert cached["target_text"] == "fuerza mayor"
    assert cached["glossary_terms"] == ["force majeure"]

    # Test miss on different language or domain
    assert mem.get("en", "fr", "legal", "Force Majeure") is None
    assert mem.get("en", "es", "medical", "Force Majeure") is None


def test_offline_deterministic_translation_glossary():
    # Legal domain
    translated, terms = translate_segment_deterministic(
        "The Force Majeure clause applies under governing law.",
        target_lang="es",
        domain="legal",
    )
    assert "fuerza mayor" in translated.lower()
    assert "ley aplicable" in translated.lower()
    assert "force majeure" in [t.lower() for t in terms]
    assert "governing law" in [t.lower() for t in terms]

    # Financial domain
    translated_fin, terms_fin = translate_segment_deterministic(
        "Quarterly EBITDA and balance sheet performance.",
        target_lang="es",
        domain="financial",
    )
    assert "balance general" in translated_fin.lower()


def test_translation_local_first_privacy_gate_f13():
    req = TranslationRequest(
        selected_text="Annual Compliance Report\nEnd-to-end data encryption at rest and in transit",
        target_language="es",
        domain="technical",
        allow_external_ai=False,  # Strict zero-egress requirement
    )
    res = execute_document_translation(req)
    assert res.provider == "local_deterministic"
    assert res.target_language == "es"
    assert res.target_language_name == "Spanish (Español)"
    assert res.total_segments >= 2
    # Check that phrase mapping or glossary worked offline
    assert "Informe Anual de Cumplimiento" in res.translated_text or "Cifrado" in res.translated_text


def test_typographic_expansion_and_font_scaling():
    req = TranslationRequest(
        selected_text="Short title",
        target_language="es",
        domain="general",
        allow_external_ai=False,
    )
    res = execute_document_translation(req)
    assert res.average_expansion_factor > 0.0
    for seg in res.segments:
        assert 0.70 <= seg.suggested_font_scale <= 1.0


def test_api_translate_endpoint(client):
    payload = {
        "selected_text": "Annual Compliance Report\nKey Controls and governing law.",
        "target_language": "es",
        "domain": "legal",
        "allow_external_ai": False,
        "scope": "selection",
    }
    resp = client.post("/api/ai/translate", json=payload)
    assert resp.status_code == 200
    data = resp.json()
    assert data["target_language"] == "es"
    assert data["provider"] == "local_deterministic"
    assert data["total_segments"] >= 2
    assert "fuerza mayor" in data["translated_text"] or "ley aplicable" in data["translated_text"] or "Informe" in data["translated_text"]


def test_api_translate_glossaries_endpoint(client):
    resp = client.get("/api/ai/translate/glossaries")
    assert resp.status_code == 200
    data = resp.json()
    assert "supported_languages" in data
    assert "legal" in data["domains"]
    assert "financial" in data["domains"]
    assert "es" in data["supported_languages"]


def test_api_ingest_multi_format_markdown(client):
    md_content = b"# Security Whitepaper\n\nCloud multi-tenancy and data isolation protocols."
    files = {"file": ("whitepaper.md", md_content, "text/markdown")}
    resp = client.post("/api/ai/ingest-multi-format?allow_external_ai=false", files=files)
    assert resp.status_code == 200
    data = resp.json()
    assert data["status"] == "success"
    assert data["detected_format"] == "MARKDOWN"
    assert data["page_count"] >= 1
    assert data["chunks_count"] >= 1
    assert "doc_id" not in data or data["document_id"].startswith("mf_")

    # Verify download of synthesized PDF
    dl_resp = client.get(f"/api/ai/documents/{data['document_id']}/download")
    assert dl_resp.status_code == 200
    assert dl_resp.headers["content-type"] == "application/pdf"
    assert dl_resp.content.startswith(b"%PDF-")


def test_api_translate_tenant_isolation_f05(client):
    # Ingest a document under a specific tenant user
    md_content = b"# Confidential M&A Brief\nFinancial terms and valuation."
    files = {"file": ("ma_brief.md", md_content, "text/markdown")}
    headers = {"X-User-Id": "corp_exec_99"}
    resp = client.post(
        "/api/ai/ingest-multi-format?allow_external_ai=false",
        files=files,
        headers=headers,
    )
    assert resp.status_code == 200
    doc_id = resp.json()["document_id"]

    # Anonymous user querying authenticated document should receive 401
    anon_resp = client.post(
        "/api/ai/translate",
        json={"document_id": doc_id, "target_language": "es"},
    )
    assert anon_resp.status_code == 401

    # Unauthorized cross-tenant user should receive 403
    cross_resp = client.post(
        "/api/ai/translate",
        json={"document_id": doc_id, "target_language": "es"},
        headers={"X-User-Id": "attacker_user_12"},
    )
    assert cross_resp.status_code == 403

    # Legitimate owner should succeed
    owner_resp = client.post(
        "/api/ai/translate",
        json={"document_id": doc_id, "target_language": "es"},
        headers={"X-User-Id": "corp_exec_99"},
    )
    assert owner_resp.status_code == 200
    assert owner_resp.json()["target_language"] == "es"
